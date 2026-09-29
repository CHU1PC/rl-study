import json
import shutil
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from itertools import pairwise
from pathlib import Path
from typing import Final

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.oxml.xmlchemy import BaseOxmlElement
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide
from pptx.text.text import TextFrame
from pptx.util import Emu, Inches, Pt

ROOT: Final = Path(__file__).resolve().parent.parent
ASSETS: Final = ROOT / "slides" / "assets"
OUT: Final = ROOT / "強化学習入門.pptx"

JP: Final = "Hiragino Sans"
CODE: Final = "Menlo"

# 文字は黒と赤 + グレー. 黒=見出しと本文と図形, 赤=強調, グレー=注記. 濃紺は見出しの下線だけに使う.
INK: Final = "000000"
RED: Final = "C00000"
SUB: Final = "595959"
NAVY: Final = "1F3864"
RULE: Final = "DDDDDD"
CODE_BG: Final = "F2F2F2"
WHITE: Final = "FFFFFF"

# フォントサイズの段. HEAD は見出し, LEAD は見出しの下の主張文, ONE は一文だけのページ用.
HEAD: Final = 40
LEAD: Final = 28
BODY: Final = 20
NOTE: Final = 16
ONE: Final = 54

# 16:9 の余白. 全要素をこの枠に収める.
LEFT: Final = 0.8
RIGHT: Final = 12.53
TOP: Final = 0.3
BOTTOM: Final = 6.75
WIDTH: Final = RIGHT - LEFT

# 見出しの下線と主張文の縦位置.
HEAD_RULE_Y: Final = 1.0
CLAIM_Y: Final = 1.12

BANNED: Final = ("最適化", "活用", "多角的", "さまざまな", "シナジー", "推進", "向上")


@dataclass(frozen=True, slots=True)
class Span:
    text: str = field()
    size: int = field(default=BODY)
    bold: bool = field(default=False)
    color: str = field(default=INK)
    font: str = field(default=JP)


type Para = list[Span]


@dataclass(frozen=True, slots=True)
class Box:
    x: float = field()
    y: float = field()
    w: float = field()
    h: float = field()


@dataclass(frozen=True, slots=True)
class Section:
    number: int = field()
    name: str = field()
    summary: str = field()  # 目次で節の名前の右に添える中身

    @property
    def heading(self) -> str:
        return f"{self.number}. {self.name}"


BASICS: Final = Section(1, "強化学習とは", "教師あり学習との違い、エージェントと環境")
Q_LEARNING: Final = Section(2, "Q学習", "問い、Q学習、探索、迷路のデモ")
POLICY_GRADIENT: Final = Section(3, "方策勾配法", "分類器として学ぶ、損失、CartPoleのデモ")
LLM: Final = Section(4, "LLMと強化学習", "SFT、RLHFとGRPO")
FIT: Final = Section(5, "向き不向き", "評価できれば学べる、向かない場面")
SUMMARY: Final = Section(6, "まとめ", "今日の一文に戻る")
SECTIONS: Final = (BASICS, Q_LEARNING, POLICY_GRADIENT, LLM, FIT, SUMMARY)


@dataclass(slots=True)
class Deck:
    prs: PresentationType = field()
    titles: list[tuple[str, int]] = field(default_factory=list[tuple[str, int]])
    texts: list[str] = field(default_factory=list[str])


def rgb(hex_color: str) -> RGBColor:
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def sp(text: str, size: int = BODY, *, bold: bool = False, color: str = INK, font: str = JP) -> Span:
    return Span(text, size, bold, color, font)


def line_weight(line: str) -> float:
    """LibreOffice で太字の 1 行が何文字分の幅になるかの見積もり.

    Args:
        line: 1 行分の文字列.

    Returns:
        全角 1 文字を 1 とした幅.
    """

    def is_full(ch: str) -> bool:
        return unicodedata.east_asian_width(ch) in ("W", "F", "A")

    full = sum(1 for ch in line if is_full(ch))
    half = len(line) - full
    boundaries = sum(1 for a, b in pairwise(line) if is_full(a) != is_full(b))
    return full + 0.55 * half + 0.95 * boundaries


def spacing_for(size: int) -> float:
    if size >= HEAD:
        return 1.0
    if size >= LEAD:
        return 1.15
    return 1.35


def fill_frame(frame: TextFrame, paras: list[Para], *, align: PP_ALIGN, gap: float) -> None:
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(0)
    for i, para in enumerate(paras):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing_for(max(s.size for s in para))
        p.space_after = Pt(gap)
        for s in para:
            run = p.add_run()
            run.text = s.text
            run.font.name = s.font
            run.font.size = Pt(s.size)
            run.font.bold = s.bold
            run.font.color.rgb = rgb(s.color)


def text(
    deck: Deck,
    slide: Slide,
    box: Box,
    paras: list[Para],
    *,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    gap: float = 6,
) -> None:
    shape = slide.shapes.add_textbox(Inches(box.x), Inches(box.y), Inches(box.w), Inches(box.h))
    shape.text_frame.vertical_anchor = anchor
    fill_frame(shape.text_frame, paras, align=align, gap=gap)
    deck.texts.extend(s.text for para in paras for s in para)


def strip_style(element: BaseOxmlElement) -> None:
    """図形の既定スタイル (影など) を消す."""
    style = element.find(qn("p:style"))
    if style is not None:
        # LibreOffice は shadow.inherit=False だけでは影を残すので, 要素ごと消す.
        element.remove(style)


def rect(
    deck: Deck,
    slide: Slide,
    box: Box,
    paras: list[Para],
    *,
    fill: str | None = None,
    line: str | None = None,
    align: PP_ALIGN = PP_ALIGN.CENTER,
) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(box.x), Inches(box.y), Inches(box.w), Inches(box.h))
    strip_style(shape.element)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(2)
    frame = shape.text_frame
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_frame(frame, paras, align=align, gap=0)
    frame.margin_left = frame.margin_right = Inches(0.15)
    deck.texts.extend(s.text for para in paras for s in para)


def code(deck: Deck, slide: Slide, box: Box, lines: list[str], *, color: str = INK) -> None:
    rect(
        deck,
        slide,
        box,
        [[sp(line, BODY, color=color, font=CODE)] for line in lines],
        fill=CODE_BG,
        align=PP_ALIGN.LEFT,
    )


def arrow(
    slide: Slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str = INK,
    head: bool = True,
    width: float | None = None,
) -> None:
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    strip_style(conn.element)
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(width if width is not None else 2.5 if head else 1)
    if head:
        for ln in conn.element.iter(qn("a:ln")):
            ln.append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle" w="med" len="med"/>'))


def picture(slide: Slide, path: Path, box: Box) -> None:
    slide.shapes.add_picture(str(path), Inches(box.x), Inches(box.y), width=Inches(box.w))


def table(
    deck: Deck,
    slide: Slide,
    x: float,
    y: float,
    widths: list[float],
    rows: list[list[str]],
    *,
    highlight: tuple[int, int],
    row_h: float = 0.62,
) -> None:
    shape = slide.shapes.add_table(
        len(rows), len(widths), Inches(x), Inches(y), Inches(sum(widths)), Inches(row_h * len(rows))
    )
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    for style_id in shape.element.iter(qn("a:tableStyleId")):
        style_id.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    for c, w in enumerate(widths):
        tbl.columns[c].width = Inches(w)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_h)
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.12)
            cell.margin_right = Inches(0.12)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(INK)
                span = sp(value, bold=True, color=WHITE)
            elif (r, c) == highlight:
                cell.fill.background()
                span = sp(value, bold=True, color=RED)
            else:
                cell.fill.background()
                span = sp(value, bold=c == 0)
            fill_frame(cell.text_frame, [[span]], align=PP_ALIGN.LEFT, gap=0)
            deck.texts.append(value)
    # 見出し以外の行の下にだけ細い罫線を引く. lnB は tcPr の中で塗りより前に置く.
    for r, tr in enumerate(shape.element.iter(qn("a:tr"))):
        if r == 0:
            continue
        for tc in tr.iter(qn("a:tc")):
            tc_pr = tc.find(qn("a:tcPr"))
            if tc_pr is None:
                continue
            tc_pr.insert(
                0,
                parse_xml(
                    f'<a:lnB {nsdecls("a")} w="12700"><a:solidFill><a:srgbClr val="{RULE}"/></a:solidFill></a:lnB>'
                ),
            )


def new_slide(deck: Deck, heading: str | None, title: str | None, notes: str) -> Slide:
    slide = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    number = len(deck.prs.slides)
    if heading is not None:
        text(deck, slide, Box(LEFT, TOP, WIDTH, 0.65), [[sp(heading, HEAD, bold=True)]], gap=0)
        arrow(slide, LEFT, HEAD_RULE_Y, RIGHT, HEAD_RULE_Y, color=NAVY, head=False, width=2.25)
        deck.titles.append((heading, HEAD))
    if title is not None:
        text(
            deck,
            slide,
            Box(LEFT, CLAIM_Y, WIDTH, 1.0),
            [[sp(line, LEAD, bold=True)] for line in title.split("\n")],
            gap=0,
        )
        deck.titles.append((title, LEAD))
    if number > 1:
        text(deck, slide, Box(RIGHT - 1.2, 6.35, 1.2, 0.4), [[sp(str(number), NOTE, color=SUB)]], align=PP_ALIGN.RIGHT)
    frame = slide.notes_slide.notes_text_frame
    if frame is not None:
        frame.text = notes
    return slide


def content_top(title: str) -> float:
    return 1.95 if "\n" not in title else 2.45


def one_sentence(deck: Deck, heading: str, lines: list[str], notes: str) -> None:
    slide = new_slide(deck, heading, None, notes)
    text(
        deck,
        slide,
        Box(LEFT, 1.5, WIDTH, 4.6),
        [[sp(line, ONE, bold=True, color=RED)] for line in lines],
        anchor=MSO_ANCHOR.MIDDLE,
        gap=0,
    )
    deck.titles.append(("\n".join(lines), ONE))


def contents(deck: Deck, current: Section | None, notes: str) -> None:
    """目次のページを作る.

    Args:
        deck: 追加先のデッキ.
        current: これから話す節. 赤にして他を灰色にする. None なら全部の節を黒で並べる.
        notes: 発表者ノート.
    """
    slide = new_slide(deck, "目次", None, notes)

    def row(i: int, name: str, summary: str, name_color: str, summary_color: str) -> None:
        y = 1.5 + i * 0.68
        text(deck, slide, Box(LEFT + 0.45, y, 4.6, 0.55), [[sp(name, LEAD, bold=True, color=name_color)]])
        text(deck, slide, Box(5.9, y + 0.07, 6.6, 0.45), [[sp(summary, color=summary_color)]])

    for i, section in enumerate(SECTIONS):
        if current is None:
            row(i, section.heading, section.summary, INK, SUB)
        elif section is current:
            text(deck, slide, Box(LEFT, 1.5 + i * 0.68, 0.45, 0.55), [[sp("▶", LEAD, color=RED)]])
            row(i, section.heading, section.summary, RED, INK)
        else:
            row(i, section.heading, section.summary, SUB, SUB)
    row(len(SECTIONS), "付録", "質疑用の5枚", SUB, SUB)


@dataclass(frozen=True, slots=True)
class CartpoleSummary:
    before: int = field()  # 学習前に倒れるまでのステップ数
    after: int = field()  # 学習後に倒れるまでのステップ数
    episodes: int = field()  # 学習に使ったエピソード数


def load_cartpole_summary() -> CartpoleSummary:
    data: dict[str, int] = json.loads((ASSETS / "cartpole_summary.json").read_text(encoding="utf-8"))
    return CartpoleSummary(data["before"], data["after"], data["episodes"])


def s01_cover(deck: Deck) -> None:
    slide = new_slide(
        deck,
        None,
        None,
        "【0:00-0:30】加藤忠です。今日は強化学習の基礎を30分で話します。"
        "教師あり学習（正解ラベルを使う学習）は知っている前提で、その違いから入ります。",
    )
    text(deck, slide, Box(LEFT, 2.2, WIDTH, 1.0), [[sp("強化学習の基礎", ONE, bold=True, color=INK)]], gap=0)
    text(deck, slide, Box(LEFT, 3.45, WIDTH, 0.6), [[sp("正解ラベルの代わりに、報酬で学ぶ", LEAD)]])
    text(deck, slide, Box(LEFT, 5.6, WIDTH, 0.5), [[sp("社内勉強会　加藤 忠", BODY, color=SUB)]])


def s02_conclusion(deck: Deck) -> None:
    one_sentence(
        deck,
        "今日の結論",
        ["強化学習は、", "報酬で重み付けした", "教師あり学習として読める"],
        "【0:30-1:15】結論を先に言います。強化学習は「報酬で重み付けした教師あり学習」として読めます。\n"
        "今日の30分は、この一文を分解します。",
    )


def s03_contents(deck: Deck) -> None:
    contents(
        deck,
        None,
        "【1:15-1:45】今日の流れです。\n"
        "1. 教師あり学習と何が違うかを押さえる。\n"
        "2. Q学習を回帰として読む。\n"
        "3. 方策勾配法を分類として読む。\n"
        "4. LLMの学習（RLHF、GRPO）も同じ式で動いていることを見る。\n"
        "5と6. 向き不向きを話して、最初の一文に戻る。\n"
        "付録は質疑用です。",
    )


def s04_difference(deck: Deck) -> None:
    title = "強化学習には正解ラベルがなく、\n報酬は遅れて届く"
    slide = new_slide(
        deck,
        BASICS.heading,
        title,
        "【1:45-3:00】違いは3つです。\n"
        "1つ目、手がかりが正解ラベルではなく、報酬という点数だけ。\n"
        "2つ目が一番大事で、報酬が遅れて届く。迷路なら、ゴールに着いた瞬間に+1がもらえるだけで、"
        "途中のどの1手が良かったのかは誰も教えてくれない。\n"
        "3つ目、データが固定ではない。自分がどう動くかで、次に見る状況が変わる。これは後で「探索」の話につながります。",
    )
    top = content_top(title)
    table(
        deck,
        slide,
        LEFT,
        top,
        [2.2, 3.6, 3.6],
        [
            ["観点", "教師あり学習", "強化学習"],
            ["手がかり", "正解ラベル", "報酬（点数）"],
            ["届くとき", "1問ごとに、すぐ", "何手も後"],
            ["データ", "固定のデータセット", "自分の行動で変わる"],
        ],
        highlight=(2, 2),
    )
    text(
        deck,
        slide,
        Box(LEFT, top + 2.8, 7.5, 1.0),
        [[sp("迷路なら、ゴールに着いた瞬間に+1がもらえるだけで、途中の1手ごとの正解は誰も教えない。")]],
    )


def s05_loop(deck: Deck) -> None:
    title = "エージェントは、報酬の合計Gが\n最大になるように行動を選ぶ"
    slide = new_slide(
        deck,
        BASICS.heading,
        title,
        "【3:00-4:45】用語を4つだけ使います。状態、行動、報酬、方策。\n"
        "方策 π(a|s) は「状態 s で行動 a を選ぶ確率」で、これを学ぶのが目標です。\n"
        "目的は、目の前の報酬ではなく、この先もらう報酬の合計 G を大きくすること。\n"
        "γ（割引率）は、先の報酬ほど軽く数えるための重みです。"
        "γ=0.9 なら、10手先の+1は0.35（0.9の10乗）として数えます。\n"
        "近道ほど G が大きくなるので、γ を入れると最短経路が好まれます。ここからは迷路を例に話します。",
    )
    top = content_top(title)
    rect(
        deck,
        slide,
        Box(1.3, top + 0.15, 3.4, 1.3),
        [[sp("エージェント", LEAD, bold=True, color=WHITE)], [sp("方策 π(a|s)", BODY, color=WHITE)]],
        fill=INK,
    )
    rect(
        deck,
        slide,
        Box(8.6, top + 0.15, 3.4, 1.3),
        [[sp("環境", LEAD, bold=True, color=RED)], [sp("迷路、ゲームなど", BODY)]],
        line=INK,
    )
    arrow(slide, 4.7, top + 0.5, 8.6, top + 0.5)
    text(deck, slide, Box(4.7, top + 0.03, 3.9, 0.45), [[sp("行動 a")]], align=PP_ALIGN.CENTER)
    arrow(slide, 8.6, top + 1.1, 4.7, top + 1.1)
    text(deck, slide, Box(4.7, top + 1.18, 3.9, 0.45), [[sp("次の状態 s′、報酬 r")]], align=PP_ALIGN.CENTER)
    text(
        deck,
        slide,
        Box(LEFT, top + 1.95, WIDTH, 0.6),
        [[sp("G = r₁ + γ·r₂ + γ²·r₃ + γ³·r₄ + …", LEAD, color=RED, font=CODE)]],
        align=PP_ALIGN.CENTER,
    )
    text(
        deck,
        slide,
        Box(LEFT, top + 2.7, 10.5, 0.9),
        [
            [sp("目的は、目の前の報酬ではなく、報酬の合計Gを大きくすること。")],
            [sp("γ（割引率）は、先の報酬ほど軽く数えるための0〜1の重み。")],
        ],
    )
    text(
        deck,
        slide,
        Box(LEFT, top + 3.75, 10.5, 0.35),
        [[sp("迷路の例：状態＝いるマス、行動＝上下左右、報酬＝ゴールで+1、落とし穴で−1、他は0", NOTE, color=SUB)]],
    )


def s06_to_q_learning(deck: Deck) -> None:
    contents(
        deck,
        Q_LEARNING,
        "【4:45-5:00】次は Q学習です。報酬がゴールにしか出ない迷路で、途中の1手をどう評価するかを見ます。",
    )


def s07_question(deck: Deck) -> None:
    one_sentence(
        deck,
        Q_LEARNING.heading,
        ["ゴールでしか報酬が出ないのに、", "最初の1手の良し悪しを", "どう知ればよいか"],
        "【5:00-5:45】問いかけです。少し考えてもらいます（15秒）。\n"
        "スタート地点では、上下左右どれを選んでも報酬は0。報酬が出るのは8手先のゴールだけ。"
        "それでも最初の1手を選べるようになるには、何を学べばよいか。答えは次の Q学習 です。",
    )


def s08_q_learning(deck: Deck) -> None:
    title = "Q学習は、1歩先の自分の予測を\n正解ラベルにした回帰である"
    slide = new_slide(
        deck,
        Q_LEARNING.heading,
        title,
        "【5:45-8:45】今日の式その1です。\n"
        "Q(s,a) は「この手を打ったら、この先いくらもらえそうか」の見込みです。"
        "学習後のスタート地点では、上0.43、右0.39、下0.48、左0.43で、下が一番大きいので下に進む。\n"
        "迷路なら25マス×4手=100個の数字の表になる。問題はこの表をどう埋めるか。\n"
        "教師ありの回帰は、予測を正解 y に少しずつ近づけます。Q学習も同じ形ですが、y を人が付けるのではなく、"
        "「1歩動いて分かった報酬 r」と「次のマスでの自分の見込み max Q(s′,a′)」から自分で作ります。\n"
        "ゴールの隣のマスは、y にゴールの+1が直接入るので正しい値になる。その隣のマスは、隣のQを y に使う。"
        "こうして報酬が1マスずつ逆向きに伝わります。これがさっきの問いの答えです。",
    )
    top = content_top(title)
    text(
        deck,
        slide,
        Box(LEFT, top, WIDTH, 0.8),
        [
            [sp("Q(s,a)", bold=True, color=RED), sp("＝マスsで手aを打った後に、この先もらえるGの見込み")],
            [
                sp("Qが分かれば、各マスでQが一番大きい手を選ぶだけでよい。", NOTE, color=SUB),
                sp("π(s) = argmaxₐ Q(s,a)", NOTE, color=SUB, font=CODE),
            ],
        ],
    )
    code(deck, slide, Box(LEFT, top + 1.05, WIDTH, 0.7), ["Q(s,a) ← Q(s,a) + α · ( r + γ · maxₐ′ Q(s′,a′) − Q(s,a) )"])
    rows = [
        ("y = r + γ · maxₐ′ Q(s′,a′)", "目標：1歩先で得た報酬＋次のマスの自分の見込み"),
        ("y − Q(s,a)", "誤差：予測の外れ（TD誤差と呼ぶ）"),
        ("α", "学習率：誤差の何割だけ寄せるか（デモでは0.5）"),
    ]
    for i, (term, meaning) in enumerate(rows):
        y = top + 1.9 + i * 0.5
        text(deck, slide, Box(LEFT, y, 4.7, 0.45), [[sp(term, color=RED, font=CODE)]])
        text(deck, slide, Box(5.7, y, 6.8, 0.45), [[sp(meaning)]])
    text(
        deck,
        slide,
        Box(LEFT, top + 3.45, WIDTH, 0.6),
        [[sp("回帰ならyは人が付ける。Q学習はyを自分で作る。", LEAD, bold=True, color=RED)]],
    )


def s09_explore(deck: Deck) -> None:
    title = "学習中は、わざと一定の割合で\nでたらめな手を打つ"
    slide = new_slide(
        deck,
        Q_LEARNING.heading,
        title,
        "【8:45-9:45】3つ目の違い「自分の行動がデータを作る」への対処です。\n"
        "いつもQが一番大きい手だけ選ぶと、最初にたまたま見つけた遠回りの道を使い続けて、近道を一度も試しません。\n"
        "そこで確率εだけランダムに動く。これを ε-greedy と呼びます。デモでは ε=0.1、10回に1回です。",
    )
    top = content_top(title)
    text(
        deck,
        slide,
        Box(LEFT, top, 5.3, 2.6),
        [
            [sp("いつも一番良い手だけ選ぶと", bold=True)],
            [sp("最初に見つけた遠回りの道を使い続け、近道を一度も試さない。")],
        ],
        gap=10,
    )
    arrow(slide, 6.65, top, 6.65, top + 2.4, color=RULE, head=False)
    text(
        deck,
        slide,
        Box(7.1, top, 5.4, 2.6),
        [
            [sp("ε-greedy", bold=True, color=RED)],
            [sp("確率εでランダムな手、残りの1−εでQが一番大きい手を選ぶ。")],
            [sp("デモでは10回に1回（ε＝0.1）。")],
        ],
        gap=10,
    )
    text(
        deck,
        slide,
        Box(LEFT, 5.3, WIDTH, 0.8),
        [
            [
                sp(
                    "まだ知らない手を試す「探索」と、知っている良い手を使う「利用」の釣り合いを取る仕組みになる。",
                    NOTE,
                    color=SUB,
                )
            ]
        ],
    )


def s10_demo_grid(deck: Deck) -> None:
    title = "報酬はゴールにしかないのに、\n価値は1マスずつ逆向きに広がる"
    slide = new_slide(
        deck,
        Q_LEARNING.heading,
        title,
        "【9:45-12:30】デモ1。ここでターミナルに切り替えて、ライブで回します（数秒で終わる）。\n"
        "  uv run python -m rl_study.gridworld\n"
        "見どころ：\n"
        "・エピソード1は200手さまよってもゴールに着かない。Qはすべて0のまま。\n"
        "・エピソード2でゴールの隣のマスだけ色が付く。3でその隣。1回に1マスずつ逆向きに広がる。\n"
        "・エピソード50では最短の8手で着く。Sの値0.48は 0.9の7乗≈0.478 と一致する（7手先の+1を割り引いた値）。\n"
        "・5回目（14手）と16回目（10手）は X に落ちて−1で終わる。10回に1回（ε＝0.1）無作為に動くため。"
        "学習が進んでも、50回目以降の約6%は X に落ちる。\n"
        "動かない場合はこのスライドの4コマで説明する。",
    )
    top = content_top(title)
    picture(slide, ASSETS / "gridworld_snapshots.png", Box(LEFT, top, WIDTH, 3.23))
    text(
        deck,
        slide,
        Box(LEFT, top + 3.33, WIDTH, 0.45),
        [
            [
                sp(
                    "濃紺ほどQの最大値が大きい。矢印は各マスでQが最大の手、オレンジの線はそのエピソードで通った道。",
                    NOTE,
                    color=SUB,
                )
            ]
        ],
    )
    text(deck, slide, Box(LEFT, top + 3.8, 8.0, 0.4), [[sp("uv run python -m rl_study.gridworld", NOTE, font=CODE)]])


def s11_to_policy(deck: Deck) -> None:
    contents(
        deck,
        POLICY_GRADIENT,
        "【12:30-12:45】次は方策勾配法です。Q を使わず、手の選び方そのものを学ぶ方法で、CartPole のデモまで進みます。",
    )


def s12_policy(deck: Deck) -> None:
    title = "方策勾配法は、Qを経由せず\n方策そのものを分類器として学ぶ"
    slide = new_slide(
        deck,
        POLICY_GRADIENT.heading,
        title,
        "【12:45-14:15】ここから後半です。"
        "マスが多くて表に書けない問題では、Q をニューラルネットで近似します（DQN、付録）。\n"
        "ここからは別の道で、Q を学んでから手を選ぶのではなく、方策 π を直接学びます。\n"
        "π_θ(a|s) は、行動をクラスとする softmax 分類器です。\n"
        "LLM も「次のトークンを選ぶ分類器」なので、この形にそのまま乗ります（最後に話します）。",
    )
    top = content_top(title)
    rows = [
        ("価値ベース（Q学習、DQN）", SUB, SUB, ["状態 s", "Q(s,a)", "Qが最大の手を選ぶ"], None),
        ("方策ベース（方策勾配法）", RED, INK, ["状態 s", "π_θ(a|s)　softmax", "確率で手を選ぶ"], INK),
    ]
    xs = [(LEFT, 2.6), (4.4, 3.8), (9.2, 3.3)]
    for i, (label, label_color, color, cells, middle_fill) in enumerate(rows):
        y = top + i * 1.55
        text(deck, slide, Box(LEFT, y, WIDTH, 0.4), [[sp(label, NOTE, bold=True, color=label_color)]])
        for j, (cell, (x, w)) in enumerate(zip(cells, xs, strict=True)):
            filled = middle_fill if j == 1 else None
            fg = WHITE if filled is not None else INK
            rect(deck, slide, Box(x, y + 0.45, w, 0.8), [[sp(cell, color=fg)]], fill=filled, line=color)
        arrow(slide, LEFT + 2.6, y + 0.85, 4.4, y + 0.85, color=color)
        arrow(slide, 8.2, y + 0.85, 9.2, y + 0.85, color=color)
    text(
        deck,
        slide,
        Box(LEFT, top + 3.25, 7.5, 0.5),
        [[sp("π_θ は、行動をクラスとするsoftmax分類器になる。")]],
    )


def s13_policy_loss(deck: Deck) -> None:
    title = "方策勾配法の損失は、\nAで重み付けしたcross-entropyになる"
    slide = new_slide(
        deck,
        POLICY_GRADIENT.heading,
        title,
        "【14:15-17:00】今日の式その2です。\n"
        "教師ありの分類は、正解ラベル a* の −log π を小さくします。\n"
        "方策勾配法は、ラベルを「自分が実際に選んだ行動 a」にして、その行動がどれだけ良かったか A で重み付けします。\n"
        "A が正なら、その行動を出やすくする。負なら出にくくする。"
        "コードでは cross_entropy に A を掛けて平均するだけです。導出は付録にあります（log の微分のトリック）。\n"
        "A の作り方です。デモの CartPole（台車の上の棒を倒さないゲーム）は、"
        "棒が立っている間、1ステップごとに+1がもらえます。"
        "G をそのまま重みにすると、G は常に正なので、どの行動も押し上げてしまう。"
        "平均を引くと、平均より長く立てたエピソードの行動だけが増えます。"
        "平均を引いても勾配の期待値は変わらず、ばらつきだけが減ります（付録）。\n"
        "advantages の1行を覚えておいてください。LLM の GRPO でほぼそのまま出てきます。",
    )
    top = content_top(title)
    rows = [
        ("教師あり", INK, "loss = −log π(a*|s)", "a* は人が付けた正解"),
        ("方策勾配法", RED, "loss = −A · log π(a|s)", "a は自分が選んだ行動"),
    ]
    for i, (label, color, formula, note) in enumerate(rows):
        y = top + i * 0.55
        text(deck, slide, Box(LEFT, y, 2.3, 0.5), [[sp(label, bold=True, color=color)]])
        text(deck, slide, Box(3.2, y, 5.0, 0.5), [[sp(formula, color=color, font=CODE)]])
        text(deck, slide, Box(8.4, y + 0.04, 4.1, 0.5), [[sp(note, NOTE, color=SUB)]])
    code(
        deck,
        slide,
        Box(LEFT, top + 1.2, WIDTH, 1.45),
        [
            "advantages = (returns - returns.mean()) / (returns.std() + 1e-8)",
            'ce = F.cross_entropy(logits, actions, reduction="none")',
            "loss = (ce * advantages).mean()",
        ],
    )
    text(
        deck,
        slide,
        Box(LEFT, top + 2.7, 8.0, 0.4),
        [[sp("rl_study/cartpole/train.py の update 関数", NOTE, color=SUB)]],
    )
    text(
        deck,
        slide,
        Box(LEFT, top + 3.1, WIDTH, 0.9),
        [
            [sp("A＞0の行動は出やすく、A＜0の行動は出にくくなる。", bold=True, color=RED)],
            [
                sp(
                    "Aは、その回のGがバッチの平均よりどれだけ良かったか。平均を引かないと、どの行動も押し上げてしまう。",
                    NOTE,
                    color=SUB,
                )
            ],
        ],
    )


def s14_demo_cartpole(deck: Deck, summary: CartpoleSummary) -> None:
    title = f"学習前は{summary.before}ステップで倒れた棒が、\n{summary.episodes}エピソードで500ステップ立つ"
    slide = new_slide(
        deck,
        POLICY_GRADIENT.heading,
        title,
        "【17:00-20:00】デモ2。GIF はスライドショーで動きます。\n"
        f"左上は重みがランダムな方策で、{summary.before}ステップで倒れる。"
        "左下は学習後で、上限の500ステップまで立ち続ける。\n"
        f"学習に使ったのは{summary.episodes}エピソード、Mac の CPU で10秒ほどです。\n"
        "  uv run python -m rl_study.cartpole --watch\n"
        "観測は台車の位置と速度、棒の角度と角速度の4つ。行動は左右に押す2択。\n"
        "右は学習曲線です。灰色が各エピソード、濃紺が直近20回の平均。\n"
        "伸びてはいるが、250エピソードあたりで一度大きく落ち込みます。"
        "方策勾配法は、1回の更新で方策が大きく変わると崩れやすい。\n"
        "これを抑える工夫が PPO で、RLHF でも使われました（付録）。",
    )
    top = content_top(title)
    text(deck, slide, Box(LEFT, top - 0.05, 3.2, 0.33), [[sp("学習前（重みはランダム）", NOTE, color=SUB)]])
    picture(slide, ASSETS / "cartpole_before.gif", Box(LEFT, top + 0.28, 2.7, 1.8))
    text(deck, slide, Box(LEFT, top + 2.15, 3.2, 0.33), [[sp("学習後", NOTE, bold=True, color=RED)]])
    picture(slide, ASSETS / "cartpole_after.gif", Box(LEFT, top + 2.48, 2.7, 1.8))
    picture(slide, ASSETS / "cartpole_curve.png", Box(4.2, top + 0.05, 7.6, 3.99))


def s15_to_llm(deck: Deck) -> None:
    contents(
        deck,
        LLM,
        "【20:00-20:15】次は LLM です。ここまでの方策勾配法の式が、RLHF と GRPO でそのまま使われていることを見ます。",
    )


def s16_llm(deck: Deck) -> None:
    title = "LLMは次のトークンを選ぶ方策で、\nSFTは教師あり学習にあたる"
    slide = new_slide(
        deck,
        LLM.heading,
        title,
        "【20:15-22:45】ここから LLM です。"
        "状態はプロンプトとここまでの出力、行動は次のトークン、方策は LLM そのもの。\n"
        "SFT は、人が書いた正解文を真似る学習で、方策勾配法の損失のページの「教師あり」の行と同じ式です。\n"
        "SFT は正解文を書ける範囲までしか上手くなりません。RL なら、自分で生成した回答を点数で重み付けするので、"
        "正解文がなくても、点数さえ付けられれば伸ばせます。",
    )
    top = content_top(title)
    pairs = [
        ("状態 s", "プロンプト＋ここまでに出したトークン"),
        ("行動 a", "次のトークン"),
        ("方策 π", "LLM（語彙の数だけクラスがあるsoftmax）"),
        ("報酬 r", "回答全体への点数（人の好み、正誤判定）"),
    ]
    for i, (term, meaning) in enumerate(pairs):
        y = top + i * 0.5
        text(deck, slide, Box(LEFT, y, 2.2, 0.45), [[sp(term, bold=True, color=RED)]])
        text(deck, slide, Box(3.2, y, 9.3, 0.45), [[sp(meaning)]])
    code(
        deck,
        slide,
        Box(LEFT, top + 2.2, WIDTH, 1.0),
        ["SFT: loss = −log π(正解文のトークン|s)", "RL : loss = −A · log π(自分で生成したトークン|s)"],
    )
    text(
        deck,
        slide,
        Box(LEFT, top + 3.35, 10.0, 0.45),
        [[sp("SFT（Supervised Fine-Tuning）は、人が書いた正解文を真似る追加学習。", NOTE, color=SUB)]],
    )


def s17_rlhf_grpo(deck: Deck) -> None:
    title = "RLHFもGRPOも同じ式で動き、\nAの作り方だけが違う"
    slide = new_slide(
        deck,
        LLM.heading,
        title,
        "【22:45-25:15】CartPole、RLHF、GRPO を並べると、違うのは報酬と A の作り方だけです。\n"
        "RLHF は、人の好みを予測するモデル（報酬モデル）で点数を付け、PPO で学習します。"
        "PPO は、1回の更新で方策を変えすぎないよう上限をかける方法で、"
        "A は状態の価値 V(s) を別のネットで予測し、G から引いて作ります（付録）。\n"
        "GRPO は、同じ問いに回答をいくつか出させ、その中で報酬を標準化して A にします。"
        "方策勾配法の損失のページにある advantages の1行と、ほぼ同じ式です。DeepSeek-R1 の学習に使われました。",
    )
    top = content_top(title)
    table(
        deck,
        slide,
        LEFT,
        top,
        [3.3, 3.9, 3.9],
        [
            ["手法", "報酬", "Aの作り方"],
            ["CartPole（デモ）", "立っていたステップ数", "バッチ内で標準化"],
            ["RLHF（PPO）", "人の好みの予測モデル", "G − V(s)"],
            ["GRPO", "正誤判定などのルール", "回答群の中で標準化"],
        ],
        highlight=(3, 2),
    )
    text(
        deck,
        slide,
        Box(LEFT, top + 2.6, 10.5, 0.35),
        [[sp("V(s)：状態sから先にもらえるGの予測。別のネットで学ぶ（PPO、付録）", NOTE, color=SUB)]],
    )
    text(
        deck,
        slide,
        Box(LEFT, top + 3.1, 7.5, 0.9),
        [[sp("GRPOは、同じ問いへの回答をいくつか出させ、平均より良かった回答のトークンを増やす。")]],
    )


def s18_to_fit(deck: Deck) -> None:
    contents(
        deck,
        FIT,
        "【25:15-25:30】次は向き不向きです。強化学習が効く場面と、遠回りになる場面を1枚ずつ話します。",
    )


def s19_value(deck: Deck) -> None:
    one_sentence(
        deck,
        FIT.heading,
        ["正解を書けなくても、", "評価できれば学習できる"],
        "【25:30-26:30】これが強化学習のいいところです。\n"
        "数学の答えが合っているか、テストが通るか、ゲームに勝ったか。どれも正解の手順は書けないし、微分もできない。"
        "でも点数は付けられる。点数さえ付けば、その点数が上がる方向に学習できる。\n"
        "正解例を真似るだけではないので、人の手本を超えることもある（AlphaGo）。",
    )


def s20_limits(deck: Deck, summary: CartpoleSummary) -> None:
    title = "正解ラベルを作れる問題では、\n強化学習は遠回りになる"
    slide = new_slide(
        deck,
        FIT.heading,
        title,
        "【26:30-28:00】向かない場面も言っておきます。\n"
        "正解ラベルを作れるなら、教師あり学習のほうが速くて安定します。\n"
        "報酬の抜け道を突かれることがある。有名な例は OpenAI が2016年に報告したボートレースのゲームで、"
        "ゴールせずに同じ場所を回ってアイテムを取り続け、点数だけを稼いだ。\n"
        f"試行回数も多い。CartPole でも{summary.episodes}エピソードかかりました。",
    )
    top = content_top(title)
    text(
        deck,
        slide,
        Box(LEFT, top, 9.5, 3.0),
        [
            [sp("・正解ラベルを作れるなら、教師あり学習が速くて安定する。")],
            [sp("・報酬の抜け道を突かれる。")],
            [
                sp(
                    "ボートレースのゲームで、ゴールせずにアイテムを取り続けて点を稼いだ例がある（OpenAI、2016年）",
                    NOTE,
                    color=SUB,
                )
            ],
            [sp(f"・試行回数が多い。CartPoleでも{summary.episodes}エピソードかかった。")],
        ],
        gap=10,
    )
    text(
        deck,
        slide,
        Box(LEFT, top + 3.2, WIDTH, 0.8),
        [
            [
                sp(
                    "今日扱わなかったもの：環境のモデルを学ぶ手法、過去のログだけで学ぶオフライン強化学習、探索の工夫",
                    NOTE,
                    color=SUB,
                )
            ]
        ],
    )


def s21_to_summary(deck: Deck) -> None:
    contents(deck, SUMMARY, "【28:00-28:15】最後にまとめです。")


def s22_summary(deck: Deck) -> None:
    title = "強化学習は、報酬で重み付けした\n教師あり学習として読める"
    slide = new_slide(
        deck,
        SUMMARY.heading,
        title,
        "【28:15-30:00】まとめです。最初の一文に戻ります。\n"
        "Q学習は、1歩先の自分の予測をラベルにした回帰。"
        "方策勾配法は、自分の行動をラベルにして A で重み付けした cross-entropy。"
        "LLM の RLHF と GRPO も同じ式で、A の作り方が違うだけ。\n"
        "コードは2つとも手元の CPU で動くので、試してみてください。質疑に移ります。",
    )
    top = content_top(title)
    text(
        deck,
        slide,
        Box(LEFT, top, 10.5, 2.9),
        [
            [sp("Q学習：", bold=True, color=RED), sp("1歩先の自分の予測をラベルにした回帰")],
            [sp("方策勾配法：", bold=True, color=RED), sp("自分の行動をラベルにし、Aで重み付けしたcross-entropy")],
            [sp("RLHF、GRPO：", bold=True, color=RED), sp("同じ式で、Aの作り方が違うだけ")],
        ],
        gap=14,
    )
    text(
        deck,
        slide,
        Box(LEFT, top + 3.1, 10.0, 0.8),
        [
            [sp("uv run python -m rl_study.gridworld", NOTE, font=CODE)],
            [sp("uv run python -m rl_study.cartpole --watch", NOTE, font=CODE)],
        ],
    )


def appendix(deck: Deck, title: str, formulas: list[str], paras: list[Para], notes: str) -> None:
    slide = new_slide(deck, "付録", title, notes)
    top = content_top(title)
    height = 0.42 * len(formulas) + 0.3
    code(deck, slide, Box(LEFT, top, WIDTH, height), formulas)
    text(deck, slide, Box(LEFT, top + height + 0.3, 10.5, 6.3 - (top + height + 0.3)), paras, gap=10)


def a01_derivation(deck: Deck) -> None:
    appendix(
        deck,
        "Q. 方策勾配の式は、どう導くのか",
        [
            "J(θ) = E_τ[ G(τ) ]",
            "∇J(θ) = Σ_τ G(τ) ∇p_θ(τ) = Σ_τ G(τ) p_θ(τ) ∇log p_θ(τ)",
            "      = E_τ[ G(τ) ∇log p_θ(τ) ]",
            "log p_θ(τ) = Σ_t log π_θ(a_t|s_t) + (θに依らない遷移の項)",
            "∇J(θ) = E_τ[ Σ_t G_t ∇log π_θ(a_t|s_t) ]",
        ],
        [
            [sp("τは1エピソードの軌跡。2行目は ∇p = p·∇log p を使っている。")],
            [sp("環境の遷移確率はθに依らないので消え、環境の中身を知らなくても勾配を計算できる。")],
            [sp("最後の式は、−G_t·log π_θ(a_t|s_t) を損失にして勾配降下したときと同じ向きになる。", NOTE, color=SUB)],
        ],
        "付録。G(τ) を時刻t以降の報酬の和 G_t に置き換えてよいのは、"
        "過去の報酬が今の行動の確率と無関係で、期待値に寄与しないため。",
    )


def a02_baseline(deck: Deck) -> None:
    appendix(
        deck,
        "Q. ベースラインを引いても、\n勾配がずれないのはなぜか",
        [
            "E_a~π[ b(s) ∇log π(a|s) ] = b(s) Σ_a π(a|s) ∇log π(a|s)",
            "                          = b(s) Σ_a ∇π(a|s)",
            "                          = b(s) ∇ Σ_a π(a|s) = b(s) ∇1 = 0",
        ],
        [
            [sp("行動に依らない b(s) を引いても、勾配の期待値は変わらない。ばらつき（分散）だけが減る。")],
            [
                sp(
                    "デモのコードは b にバッチ平均を使う。"
                    "同じエピソードのGも平均に含むので厳密には少し偏るが、実用上よく使う形。",
                    NOTE,
                    color=SUB,
                )
            ],
        ],
        "付録。確率の総和が1で一定なので、その勾配は0になる、という1行の議論。",
    )


def a03_dqn(deck: Deck) -> None:
    appendix(
        deck,
        "Q. DQNは、Q学習を\nどう大きな問題に広げたのか",
        ["y = r + γ · maxₐ′ Q_θ⁻(s′,a′)      loss = ( y − Q_θ(s,a) )²"],
        [
            [sp("Qの表をニューラルネットに置き換え、状態から手ごとのQを出す。")],
            [
                sp("経験再生：", bold=True, color=RED),
                sp("過去の遷移を貯めて何度も使う。無作為に取り出して相関を切る。"),
            ],
            [
                sp("目標ネットワーク：", bold=True, color=RED),
                sp("yを作る重みθ⁻を一定間隔でしか更新せず、目標を止める。"),
            ],
            [sp("2015年に、Atariのゲームを画面の画素だけから学習した。", NOTE, color=SUB)],
        ],
        "付録。迷路は100個の表で済みますが、ゲームの画面を状態にすると組み合わせが多すぎて表に書けません。"
        "そこで Q をニューラルネットで近似します。これが DQN（Deep Q-Network）です。\n"
        "目標 y の作り方は Q学習と同じで、損失は二乗誤差。ここで完全に回帰の形になります。\n"
        "経験再生と目標ネットワークは、どちらも教師ありの回帰に近い条件"
        "（データがばらばら、ラベルが止まっている）を作るための工夫です。",
    )


def a04_ppo(deck: Deck) -> None:
    appendix(
        deck,
        "Q. PPOは、何をしているのか",
        [
            "ρ = π_θ(a|s) / π_old(a|s)",
            "L = E[ min( ρ·A, clip(ρ, 1−ε, 1+ε)·A ) ]   を最大化する",
        ],
        [
            [sp("確率の比ρが1±εの外に出た分は、学習に使わない。εは0.2などにする。")],
            [sp("A＞0なら上げすぎる得がなく、A＜0なら下げすぎる得がない。だから同じデータで数回更新できる。")],
            [sp("Aは、状態の価値V(s)を別のネット（Critic）で予測し、Gから引いて作る。これをActor-Criticと呼ぶ。")],
            [
                sp(
                    "InstructGPT（ChatGPTの前身）の、人のフィードバックによる強化学習（RLHF）で使われた。",
                    NOTE,
                    color=SUB,
                )
            ],
        ],
        "付録。PPO（Proximal Policy Optimization）は、今の実用上の標準です。"
        "方策勾配法は、1回の更新で方策が大きく変わると崩れやすい。"
        "PPO は新しい方策と古い方策で行動の確率の比をとり、1±ε の外に出た分を学習に使わないので、1回で動きすぎない。\n"
        "ρ=1 のとき ρ·A の勾配は A·∇log π になり、方策勾配法の式に戻る。",
    )


def a05_grpo(deck: Deck) -> None:
    appendix(
        deck,
        "Q. GRPOは、PPOと何が違うのか",
        ["Â_i = ( r_i − mean(r_1, …, r_G) ) / std(r_1, …, r_G)"],
        [
            [sp("同じプロンプトへのG個の回答の中で報酬を標準化し、それをAにする。")],
            [sp("価値関数（Critic）を持たないので、メモリと計算が軽い。")],
            [sp("KLペナルティで、元のモデルから離れすぎないようにする。")],
            [sp("clipの式はPPOと同じ。", NOTE, color=SUB)],
        ],
        "付録。DeepSeekMath で提案され、DeepSeek-R1 の学習に使われた。報酬には数学の正誤判定などのルールを使える。",
    )


def check(deck: Deck) -> list[str]:
    """見出しと主張文の折り返しと禁止語を検査し, 問題の一覧を返す.

    Args:
        deck: 組み立て終えたデッキ.

    Returns:
        問題の説明. 空なら合格.
    """
    max_lines = {HEAD: 1, LEAD: 2}
    problems: list[str] = []
    for title, size in deck.titles:
        limit = WIDTH * 72 / size
        lines = title.split("\n")
        if size in max_lines and len(lines) > max_lines[size]:
            problems.append(f"{max_lines[size]}行を超える: {title!r}")
        problems.extend(
            f"1行に収まらない ({line_weight(line):.1f} > {limit:.1f}): {line!r}"
            for line in lines
            if line_weight(line) > limit
        )
    body = "".join(deck.texts).replace("強化学習", "")
    problems.extend(f"禁止語: {word}" for word in BANNED if word in body)
    return problems


def main() -> None:
    summary = load_cartpole_summary()
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    deck = Deck(prs)

    s01_cover(deck)
    s02_conclusion(deck)
    s03_contents(deck)
    s04_difference(deck)
    s05_loop(deck)
    s06_to_q_learning(deck)
    s07_question(deck)
    s08_q_learning(deck)
    s09_explore(deck)
    s10_demo_grid(deck)
    s11_to_policy(deck)
    s12_policy(deck)
    s13_policy_loss(deck)
    s14_demo_cartpole(deck, summary)
    s15_to_llm(deck)
    s16_llm(deck)
    s17_rlhf_grpo(deck)
    s18_to_fit(deck)
    s19_value(deck)
    s20_limits(deck, summary)
    s21_to_summary(deck)
    s22_summary(deck)
    a01_derivation(deck)
    a02_baseline(deck)
    a03_dqn(deck)
    a04_ppo(deck)
    a05_grpo(deck)

    # 和文フォントは <a:ea> に書かないと効かないので, 全 run の latin の後ろに ea と cs を足す.
    for slide in prs.slides:
        for latin in list(slide.element.iter(qn("a:latin"))):
            for tag in ("a:cs", "a:ea"):
                latin.addnext(parse_xml(f'<{tag} {nsdecls("a")} typeface="{JP}"/>'))

    problems = check(deck)
    for problem in problems:
        print(problem)
    if problems:
        raise SystemExit(1)

    if OUT.exists():
        backup = ROOT / "slides" / "_backup" / f"{OUT.stem}-{datetime.now():%Y%m%d-%H%M%S}.pptx"
        backup.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(OUT, backup)
        print(f"既存のファイルを退避した: {backup}")
    prs.save(str(OUT))
    print(f"{len(prs.slides)} 枚を書き出した: {OUT}")


if __name__ == "__main__":
    main()
